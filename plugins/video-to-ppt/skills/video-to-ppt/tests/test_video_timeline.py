import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


@unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'), 'FFmpeg required')
class VideoTimelineTest(unittest.TestCase):
    script = Path(__file__).resolve().parents[1]/'scripts/auto_pdf.py'

    def run_pipeline(self, video, output, *flags):
        result = subprocess.run([sys.executable, '-B', str(self.script), str(video), str(output), *flags],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads((output/'run.json').read_text())

    def test_longer_audio_does_not_extend_grid_and_generated_manifest_can_be_reused(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); video = root/'audio-longer.mp4'
            subprocess.run(['ffmpeg', '-nostdin', '-v', 'error', '-f', 'lavfi', '-i',
                'color=white:s=192x108:r=10:d=16', '-f', 'lavfi', '-i',
                'sine=frequency=400:duration=21', '-c:v', 'mpeg4', '-c:a', 'aac', str(video)],
                check=True, capture_output=True)
            first = self.run_pipeline(video, root/'first')
            self.assertAlmostEqual(first['duration_seconds'], 16)
            self.assertEqual(first['duration_basis'], 'video_stream_duration')
            self.assertEqual(first['sample_count'], 4)
            reused = self.run_pipeline(video, root/'reused', '--samples-manifest', str(root/'first/samples.json'))
            self.assertTrue(reused['sampling_reused'])
            self.assertTrue(reused['source_video_reverified'])
            self.assertEqual((reused['duration_seconds'], reused['sample_count'], reused['retained_pages']),
                             (first['duration_seconds'], 4, first['retained_pages']))
            for output in [root/'first', root/'reused']:
                samples = json.loads((output/'samples.json').read_text())['samples']
                self.assertEqual([s['timestamp_seconds'] for s in samples], [0, 5, 10, 15])

    def test_missing_stream_duration_uses_video_packets_and_explicit_one_second_grid(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); video = root/'video.mkv'
            subprocess.run(['ffmpeg', '-nostdin', '-v', 'error', '-f', 'lavfi', '-i',
                'color=white:s=192x108:r=10:d=12', '-c:v', 'ffv1', str(video)],
                check=True, capture_output=True)
            run = self.run_pipeline(video, root/'first', '--interval', '1')
            self.assertEqual(run['duration_basis'], 'video_packet_pts_range')
            self.assertAlmostEqual(run['duration_seconds'], 12)
            self.assertEqual(run['sample_count'], 12)
            self.assertEqual(run['stability_eligible_samples'], 10)
            pages = json.loads((root/'first/pages.json').read_text())['pages']
            self.assertEqual([p['timestamp_seconds'] for p in pages], [1])

    def test_rotated_video_uses_actual_sample_geometry_for_first_and_reused_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); plain = root/'plain.mp4'; rotated = root/'rotated.mp4'
            subprocess.run(['ffmpeg', '-nostdin', '-v', 'error', '-f', 'lavfi', '-i',
                'color=white:s=192x108:r=10:d=16', '-vf',
                'drawbox=x=10:y=20:w=60:h=40:color=red:t=fill', '-c:v', 'mpeg4', str(plain)],
                check=True, capture_output=True)
            subprocess.run(['ffmpeg', '-nostdin', '-v', 'error', '-display_rotation:v:0', '90',
                '-i', str(plain), '-c', 'copy', str(rotated)],
                check=True, capture_output=True)
            first = self.run_pipeline(rotated, root/'first', '--format', 'both')
            reused = self.run_pipeline(rotated, root/'reused', '--format', 'both',
                                       '--samples-manifest', str(root/'first/samples.json'))
            for output, run in [(root/'first', first), (root/'reused', reused)]:
                self.assertEqual((run['width'], run['height']), (108, 192))
                pages = json.loads((output/'pages.json').read_text())['pages']
                with Image.open(pages[0]['image']) as image:
                    self.assertEqual(image.size, (108, 192))
                pdf = PdfReader(output/'slides.pdf')
                for page in pdf.pages:
                    ratio = float(page.mediabox.width)/(float(page.mediabox.height)-20)
                    self.assertAlmostEqual(ratio, 108/192, places=6)
                presentation = Presentation(output/'slides.pptx')
                for slide in presentation.slides:
                    image = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
                    self.assertAlmostEqual(image.width/image.height, 108/192, places=6)
                self.assertTrue(run['validation']['exports']['pdf']['ordered_source_pixels_preserved'])
                self.assertTrue(run['validation']['exports']['pptx']['ordered_source_pixels_preserved'])
            self.assertEqual(first['sample_count'], reused['sample_count'])
            self.assertEqual(first['retained_pages'], reused['retained_pages'])


if __name__ == '__main__':
    unittest.main()
