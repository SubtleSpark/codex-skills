# /// script
# requires-python = ">=3.12"
# dependencies = ["numpy>=1.24", "Pillow>=10", "reportlab>=4", "pypdf>=4", "scikit-image==0.26.0", "python-pptx>=1"]
# ///
"""Fixed sampling -> original-neighbour stability -> dual hash -> SSIM -> first-qualified PDF/PPTX."""
import argparse
import io
from datetime import datetime, timezone
from functools import lru_cache
import hashlib
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time

try:
    import numpy as np
    import PIL
    from PIL import Image
    import pypdf
    from reportlab import Version as reportlab_version
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
except ImportError as error:
    raise SystemExit(f'Missing dependency {error.name}; use a Python environment with '
                     'numpy, Pillow, reportlab and pypdf installed.') from error

HASH_BITS = 64
DEFAULT_DISTANCE = 1
DEFAULT_DHASH_DISTANCE = 2
DEFAULT_HASH_MODE = 'phash-dhash'


def positive(value):
    value = float(value)
    if not math.isfinite(value) or value <= 0:
        raise argparse.ArgumentTypeError('must be finite and positive')
    return value


def distance_limit(value):
    value = int(value)
    if not 0 <= value <= HASH_BITS:
        raise argparse.ArgumentTypeError('must be an integer from 0 to 64')
    return value


@lru_cache(maxsize=1)
def dct_basis():
    # DCT-II, without a normalization factor (a uniform factor does not affect the median).
    return np.cos(np.pi * np.arange(32)[:, None] * (2 * np.arange(32)[None, :] + 1) / 64)


def phash(image):
    pixels = np.asarray(image.convert('L').resize((32, 32), Image.Resampling.LANCZOS), dtype=np.float64)
    low = (dct_basis() @ pixels @ dct_basis().T)[:8, :8].copy()
    # Remove floating point roundoff around mathematically zero coefficients.
    low[np.abs(low) < 1e-9] = 0
    bits = low.ravel() > np.median(low)
    value = 0
    for bit in bits:
        value = (value << 1) | int(bit)
    return value


def hamming(left, right):
    return (left ^ right).bit_count()


def dhash(image):
    """64 horizontal gradients: grayscale 9x8 LANCZOS, right pixel > left pixel."""
    pixels = np.asarray(image.convert('L').resize((9, 8), Image.Resampling.LANCZOS))
    value = 0
    for bit in (pixels[:, 1:] > pixels[:, :-1]).ravel():
        value = (value << 1) | int(bit)
    return value


def classify_dual_hashes(phashes, dhashes, phash_threshold, dhash_threshold):
    """Both inclusive thresholds must match the same retained representative; no chaining."""
    if len(phashes) != len(dhashes):
        raise ValueError('pHash and dHash counts must match')
    for threshold in (phash_threshold, dhash_threshold):
        if type(threshold) is not int or not 0 <= threshold <= HASH_BITS:
            raise ValueError('threshold must be an integer from 0 to 64')
    kept, assignments = [], []
    for index, value in enumerate(phashes):
        for representative in kept:
            pdistance = hamming(value, phashes[representative])
            if pdistance <= phash_threshold:
                ddistance = hamming(dhashes[index], dhashes[representative])
                if ddistance <= dhash_threshold:
                    assignments.append((representative, pdistance, ddistance))
                    break
        else:
            kept.append(index)
            assignments.append((index, 0, 0))
    return kept, assignments


def compute_hashes(paths, mode=DEFAULT_HASH_MODE):
    if mode not in ('phash', 'phash-dhash'):
        raise ValueError('unknown hash mode')
    phashes, dhashes = [], []
    timings = {'decode_grayscale_seconds': 0, 'phash_seconds': 0, 'dhash_seconds': 0}
    started = time.perf_counter()
    size = None
    for path in paths:
        tick = time.perf_counter()
        with Image.open(path) as image:
            if size is not None and image.size != size:
                raise ValueError('sample dimensions differ')
            size = image.size
            gray = image.convert('L')
        timings['decode_grayscale_seconds'] += time.perf_counter() - tick
        tick = time.perf_counter()
        phashes.append(phash(gray))
        timings['phash_seconds'] += time.perf_counter() - tick
        if mode == 'phash-dhash':
            tick = time.perf_counter()
            dhashes.append(dhash(gray))
            timings['dhash_seconds'] += time.perf_counter() - tick
    timings['total_seconds'] = time.perf_counter() - started
    return phashes, dhashes, timings


def frame_records(paths, phashes, assignments, interval, dhashes=None):
    """Keep sample indices, first representatives and both distances together."""
    if (len(paths) != len(phashes) or len(paths) != len(assignments)
            or (dhashes is not None and len(paths) != len(dhashes))):
        raise ValueError('frame, hash and assignment counts must match')
    kept = [i for i, assignment in enumerate(assignments) if assignment[0] == i]
    page_numbers = {sample: number for number, sample in enumerate(kept, 1)}
    samples = []
    for i, (path, assignment) in enumerate(zip(paths, assignments)):
        representative, distance = assignment[:2]
        item = {'sample_id': i, 'timestamp_seconds': i * interval,
                'image': str(path), 'phash': f'{phashes[i]:016x}',
                'retained': representative == i, 'representative_sample_id': representative,
                'page': page_numbers[representative], 'distance_to_representative': distance}
        if dhashes is not None:
            item.update(dhash=f'{dhashes[i]:016x}',
                        dhash_distance_to_representative=hamming(dhashes[i], dhashes[representative]))
        samples.append(item)
    pages = []
    for i in kept:
        page = {'page': page_numbers[i], 'sample_id': i, 'timestamp_seconds': samples[i]['timestamp_seconds'],
                'image': str(paths[i]), 'image_sha256': sha256(paths[i]), 'phash': f'{phashes[i]:016x}',
                'occurrences': []}
        if dhashes is not None:
            page['dhash'] = f'{dhashes[i]:016x}'
        pages.append(page)
    for sample in samples:
        occurrence = {'sample_id': sample['sample_id'], 'timestamp_seconds': sample['timestamp_seconds'],
                      'distance': sample['distance_to_representative']}
        if dhashes is not None:
            occurrence['dhash_distance'] = sample['dhash_distance_to_representative']
        pages[sample['page'] - 1]['occurrences'].append(occurrence)
    return samples, pages


def compose_stages(paths, phashes, hash_samples, hash_pages, ssim_assignments, interval, dhashes=None):
    """Preserve hash provenance while composing raw -> hash representative -> SSIM representative."""
    if len(hash_pages) != len(ssim_assignments):
        raise ValueError('SSIM assignments must cover every hash representative')
    by_hash_sample = {}
    for i, (representative, score) in enumerate(ssim_assignments):
        if type(representative) is not int or not 0 <= representative <= i or ssim_assignments[representative][0] != representative:
            raise ValueError('SSIM assignment must reference an earlier retained representative')
        by_hash_sample[hash_pages[i]['sample_id']] = (hash_pages[representative]['sample_id'], score)
    assignments = []
    for i, sample in enumerate(hash_samples):
        representative, _ = by_hash_sample[sample['representative_sample_id']]
        assignments.append((representative, hamming(phashes[i], phashes[representative])))
    samples, pages = frame_records(paths, phashes, assignments, interval, dhashes)
    for sample, old in zip(samples, hash_samples):
        representative, score = by_hash_sample[old['representative_sample_id']]
        sample.update(hash_retained=old['retained'], hash_representative_sample_id=old['representative_sample_id'],
            hash_page=old['page'], phash_distance_to_hash_representative=old['distance_to_representative'],
            ssim_input_sample_id=old['representative_sample_id'], ssim_representative_sample_id=representative,
            ssim_score_hash_rep_to_final_rep=score)
        if dhashes is not None:
            sample['dhash_distance_to_hash_representative'] = old['dhash_distance_to_representative']
    for page in pages:
        for occurrence in page['occurrences']:
            sample = samples[occurrence['sample_id']]
            occurrence.update(hash_representative_sample_id=sample['hash_representative_sample_id'],
                ssim_score_hash_rep_to_final_rep=sample['ssim_score_hash_rep_to_final_rep'])
    decisions = [{'hash_page': i+1, 'sample_id': p['sample_id'],
                  'timestamp_seconds': p['timestamp_seconds'],
                  'representative_sample_id': by_hash_sample[p['sample_id']][0],
                  'representative_timestamp_seconds': samples[by_hash_sample[p['sample_id']][0]]['timestamp_seconds'],
                  'ssim': by_hash_sample[p['sample_id']][1],
                  'retained': p['sample_id'] == by_hash_sample[p['sample_id']][0],
                  'final_page': samples[p['sample_id']]['page']}
                 for i, p in enumerate(hash_pages)]
    return samples, pages, decisions


def classify_hashes(hashes, threshold):
    """Compare only with retained representatives, in first-seen order; never chain discarded frames."""
    if type(threshold) is not int or not 0 <= threshold <= HASH_BITS:
        raise ValueError('threshold must be an integer from 0 to 64')
    kept, assignments = [], []
    for index, value in enumerate(hashes):
        for representative in kept:
            distance = hamming(value, hashes[representative])
            if distance <= threshold:
                assignments.append((representative, distance))
                break
        else:
            kept.append(index)
            assignments.append((index, 0))
    return kept, assignments


def sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def save_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def clock(seconds):
    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f'{int(hours):02d}:{int(minutes):02d}:{seconds:06.3f}'


def load_samples(manifest):
    data = json.loads(manifest.read_text())
    interval, duration, rows = data['interval_seconds'], data['duration_seconds'], data['samples']
    if (type(interval) not in (int, float) or not math.isfinite(interval) or interval <= 0
            or type(duration) not in (int, float) or not math.isfinite(duration) or duration <= 0):
        raise ValueError('invalid fixed sampling interval or duration')
    if (data.get('sampling') != 'all fixed-grid samples; no MAE prefilter'
            or len(rows) != data['sample_count'] or len(rows) != math.ceil(duration / interval)):
        raise ValueError('input must contain every fixed-grid sample, not filtered candidates')
    paths = []
    for index, row in enumerate(rows):
        if (type(row['sample_id']) is not int or row['sample_id'] != index
                or type(row['timestamp_seconds']) not in (int, float)
                or not math.isclose(row['timestamp_seconds'], index * interval, rel_tol=0, abs_tol=1e-6)):
            raise ValueError('samples must be in contiguous first-seen fixed-grid order')
        path = Path(row['image'])
        path = path if path.is_absolute() else manifest.parent / path
        if not path.is_file():
            raise ValueError(f'missing source image: {path}')
        paths.append(path.resolve())
    return data, paths


def video_duration(video, stream):
    """Use the video stream, never a longer audio/container timeline.

    Some containers (notably Matroska) omit stream duration. In that case
    inspect video packet timestamps without decoding frames or reading payloads.
    """
    try:
        duration = float(stream['duration'])
        if math.isfinite(duration) and duration > 0:
            return duration, 'video_stream_duration'
    except (KeyError, TypeError, ValueError):
        pass
    command = ['ffprobe', '-v', 'error', '-select_streams', 'v:0',
        '-show_entries', 'packet=pts_time,duration_time', '-of', 'compact=p=0:nk=0', str(video)]
    start, end = math.inf, -math.inf
    with tempfile.TemporaryFile(mode='w+') as errors, subprocess.Popen(
            command, stdout=subprocess.PIPE, stderr=errors, text=True) as process:
        for line in process.stdout:
            fields = dict(part.split('=', 1) for part in line.strip().split('|') if '=' in part)
            try:
                pts = float(fields['pts_time'])
                length = float(fields['duration_time'])
            except (KeyError, ValueError):
                continue
            if math.isfinite(pts) and math.isfinite(length) and length > 0:
                start, end = min(start, pts), max(end, pts+length)
        if process.wait():
            errors.seek(0)
            raise ValueError(f'cannot probe video packet timeline: {errors.read(4096).strip()}')
    duration = end-start
    if not math.isfinite(duration) or duration <= 0:
        raise ValueError('cannot determine video stream duration from stream metadata or packet timestamps')
    return duration, 'video_packet_pts_range'


def verify_grid_count(paths, duration, interval):
    expected = math.ceil(duration/interval)
    if len(paths) != expected:
        raise ValueError(f'incomplete fixed video grid: expected {expected} samples, got {len(paths)}')


def export_pptx(path, pages, width, height, title):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    image_height = int(presentation.slide_width * height / width)
    presentation.slide_height = image_height + Inches(.28)
    presentation.core_properties.title = title
    for page in pages:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(page['image'], 0, 0, width=presentation.slide_width, height=image_height)
        box = slide.shapes.add_textbox(Inches(.16), image_height, presentation.slide_width, Inches(.28))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = f"{page['page']:03d} / {len(pages):03d}    First qualified frame {clock(page['timestamp_seconds'])}    Sample {page['sample_id']:06d}"
        paragraph.font.size = Pt(8)
    presentation.save(path)


def verify_pptx(path, pages):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(path)
    if len(presentation.slides) != len(pages):
        raise ValueError('PPTX slide count differs from retained frames')
    for slide, page in zip(presentation.slides, pages):
        images = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if len(images) != 1:
            raise ValueError('PPTX slide must contain exactly one source image')
        with Image.open(io.BytesIO(images[0].image.blob)) as embedded, Image.open(page['image']) as source:
            if embedded.size != source.size or embedded.convert('RGB').tobytes() != source.convert('RGB').tobytes():
                raise ValueError('PPTX slide differs from its ordered first-seen image')
    return {'page_count': len(pages), 'ordered_source_pixels_preserved': True}


def export_pdf(path, pages, width, height, title, source_url):
    page_width = 960
    image_height = page_width * height / width
    pdf = canvas.Canvas(str(path), pagesize=(page_width, image_height + 20),
                        pageCompression=1, invariant=1)
    pdf.setTitle(title)
    pdf.setSubject(f'Automatic first-seen video frames; source: {source_url}')
    for page in pages:
        pdf.drawImage(ImageReader(page['image']), 0, 20, width=page_width, height=image_height)
        pdf.setFillColorRGB(.25, .25, .25)
        pdf.setFont('Helvetica', 8)
        pdf.drawString(12, 7, f"{page['page']:03d} / {len(pages):03d}    First qualified frame {clock(page['timestamp_seconds'])}    Sample {page['sample_id']:06d}")
        key = f"page-{page['page']}"
        pdf.bookmarkPage(key)
        pdf.addOutlineEntry(f"{page['page']:03d} @ {clock(page['timestamp_seconds'])}", key)
        pdf.showPage()
    pdf.save()


def verify_pdf(path, pages):
    reader = pypdf.PdfReader(path)
    if len(reader.pages) != len(pages):
        raise ValueError('PDF page count differs from retained frames')
    for number, (pdf_page, page) in enumerate(zip(reader.pages, pages), 1):
        images = list(pdf_page.images)
        if len(images) != 1:
            raise ValueError(f'PDF page {number} must have exactly one source image')
        embedded = images[0].image.convert('RGB')
        with Image.open(page['image']) as image:
            source = image.convert('RGB')
        if embedded.size != source.size or embedded.tobytes() != source.tobytes():
            raise ValueError(f'PDF page {number} differs from its first-seen source image')
    return {'page_count': len(pages), 'ordered_source_pixels_preserved': True}


def main():
    from ssim_audit import audit_ssim, require_ssim
    from stability_filter import audit_stability, restore_source_ids, restore_decisions
    started_all = time.perf_counter()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('video', type=Path)
    parser.add_argument('output', type=Path, help='new directory; existing paths are refused')
    parser.add_argument('--samples-manifest', type=Path, help='reuse all existing fixed-grid images instead of decoding video')
    parser.add_argument('--interval', type=positive, default=5, help='fixed sample interval in seconds (default: 5)')
    parser.add_argument('--hash-distance', type=distance_limit, default=DEFAULT_DISTANCE, help='pHash inclusive Hamming threshold (default: 1)')
    parser.add_argument('--hash-mode', choices=('phash', 'phash-dhash'), default=DEFAULT_HASH_MODE)
    parser.add_argument('--dhash-distance', type=distance_limit, default=DEFAULT_DHASH_DISTANCE, help='dHash inclusive threshold (default: 2)')
    parser.add_argument('--stability-threshold', type=float, default=.995,
                        help='both original adjacent SSIM scores must be >= threshold (default: .995)')
    parser.add_argument('--skip-stability', action='store_true',
                        help='explicitly bypass neighbour stability filtering; record the bypass')
    parser.add_argument('--skip-ssim', '--skip-ssim-audit', dest='skip_ssim_audit', action='store_true',
                        help='explicitly omit the entire second SSIM stage')
    parser.add_argument('--ssim-mode', choices=('dedup', 'audit'), default='dedup',
                        help='dedup removes later matching hash representatives; audit preserves stage 1')
    parser.add_argument('--ssim-threshold', type=float, default=.995, help='merge if SSIM >= threshold (default: .995)')
    parser.add_argument('--ssim-width', type=int, default=960)
    parser.add_argument('--ssim-height', type=int, default=540)
    parser.add_argument('--ssim-review-threshold', type=float, default=.995, help='report high-score pairs only; never deletes pages')
    parser.add_argument('--format', dest='output_format', choices=('pdf', 'pptx', 'both'), default='pdf')
    parser.add_argument('--title', help='document metadata title; no new cover is created')
    parser.add_argument('--source-url', default='', help='optional public source URL for provenance')
    args = parser.parse_args()
    if not args.video.is_file():
        parser.error('video does not exist')
    if args.output.exists():
        parser.error('output exists; choose a new directory')
    if (args.ssim_width < 11 or args.ssim_height < 11 or not 0 <= args.ssim_review_threshold <= 1
            or not 0 <= args.ssim_threshold <= 1 or not 0 <= args.stability_threshold <= 1):
        parser.error('SSIM dimensions must be >=11 and thresholds must be finite and in [0,1]')
    if not args.skip_ssim_audit or not args.skip_stability:
        try:
            require_ssim()
        except ValueError as error:
            parser.error(str(error) + ' Neighbour stability also requires SSIM; bypassing both '
                         'requires explicit --skip-stability --skip-ssim.')
    if args.output_format in ('pptx', 'both'):
        try:
            import pptx
        except ImportError:
            parser.error('PPTX export requires python-pptx in this Python environment')
    video, output = args.video.resolve(), args.output.resolve()
    source_hash = sha256(video)
    provenance = {}
    if args.samples_manifest:
        manifest = args.samples_manifest.resolve()
        original, paths = load_samples(manifest)
        if original['interval_seconds'] != args.interval:
            parser.error('manifest interval differs; set --interval to the existing sampling interval')
        if not original.get('source_video_sha256'):
            parser.error('sample manifest must include source_video_sha256; unverified provenance cannot be reused')
        if original['source_video_sha256'] != source_hash:
            parser.error('source video SHA256 differs from the sample manifest')
        duration = original['duration_seconds']
        command = original.get('sampling_command', [])
        provenance = {'source_manifest': str(manifest), 'source_manifest_sha256': sha256(manifest),
                      'sampling_reused': True, 'source_video_reverified': True,
                      'duration_basis': original.get('duration_basis', 'verified_manifest_fixed_grid')}
        output.mkdir(parents=True)
    else:
        if not all(shutil.which(tool) for tool in ('ffmpeg', 'ffprobe')):
            parser.error('ffmpeg and ffprobe are required for new sampling; existing samples can use --samples-manifest')
        probe = json.loads(subprocess.run(['ffprobe', '-v', 'error', '-select_streams', 'v:0',
            '-show_entries', 'stream=duration', '-of', 'json', str(video)],
            check=True, capture_output=True, text=True).stdout)
        if not probe.get('streams'):
            parser.error('input has no video stream')
        duration, duration_basis = video_duration(video, probe['streams'][0])
        output.mkdir(parents=True)
        samples_dir = output/'samples'
        samples_dir.mkdir()
        command = ['ffmpeg', '-nostdin', '-v', 'error', '-i', str(video), '-map', '0:v:0',
            '-vf', f'setpts=PTS-STARTPTS,fps=1/{args.interval}:start_time=0:round=up', '-an',
            '-c:v', 'png', '-compression_level', '2', '-start_number', '0',
            str(samples_dir/'sample-%06d.png')]
        print(json.dumps({'stage': 'fixed_interval_sampling', 'interval_seconds': args.interval}), flush=True)
        with (output/'extraction.log').open('w') as log:
            subprocess.run(command, check=True, stderr=log)
        paths = sorted(samples_dir.glob('sample-*.png'))
        provenance = {'sampling_reused': False, 'source_video_reverified': True,
                      'duration_basis': duration_basis}
    if not paths:
        raise ValueError('no sampled frames')
    verify_grid_count(paths, duration, args.interval)
    # FFmpeg may apply display rotation, so encoded video dimensions are not the output geometry.
    with Image.open(paths[0]) as first:
        width, height = first.size
    all_paths = paths
    print(json.dumps({'stage': 'original_neighbour_stability', 'samples': len(all_paths),
                      'threshold': args.stability_threshold, 'skipped': args.skip_stability}), flush=True)
    stability = audit_stability(output, all_paths, args.interval, args.ssim_width, args.ssim_height,
                                args.stability_threshold, args.skip_stability)
    eligible_ids = stability['eligible_sample_ids']
    if not eligible_ids:
        raise ValueError(stability['diagnostic'] + ' See stability.json; no fallback or export was performed')
    paths = [all_paths[i] for i in eligible_ids]
    print(json.dumps({'stage': 'stability_complete', 'sample_count': len(all_paths),
                      'eligible_count': len(paths), 'rejected_count': len(all_paths)-len(paths),
                      'timings': stability['timings']}), flush=True)
    print(json.dumps({'stage': 'hash_filter', 'samples': len(paths), 'mode': args.hash_mode,
                      'phash_threshold': args.hash_distance, 'dhash_threshold': args.dhash_distance}), flush=True)
    hashes, dhashes, hash_timings = compute_hashes(paths, args.hash_mode)
    started = time.perf_counter()
    if args.hash_mode == 'phash-dhash':
        kept, assignments = classify_dual_hashes(hashes, dhashes, args.hash_distance, args.dhash_distance)
    else:
        kept, assignments = classify_hashes(hashes, args.hash_distance)
    comparison_seconds = time.perf_counter()-started
    samples, pages = frame_records(paths, hashes, assignments, args.interval,
                                   dhashes if args.hash_mode == 'phash-dhash' else None)
    print(json.dumps({'stage': 'hash_filter_complete', 'retained_pages': len(pages),
                      'hash_seconds': hash_timings, 'comparison_seconds': comparison_seconds}), flush=True)
    plugin_manifest = Path(__file__).resolve().parents[3]/'.codex-plugin/plugin.json'
    version = json.loads(plugin_manifest.read_text())['version'] if plugin_manifest.is_file() else 'unpackaged'
    metadata = {'pipeline': 'fixed_interval_stability_hash_then_ssim_first_qualified', 'plugin_version': version,
        'entrypoint': str(Path(__file__).resolve()), 'entrypoint_sha256': sha256(Path(__file__).resolve()),
        'source_video': str(video), 'source_video_sha256': source_hash,
        'source_url': args.source_url or (original.get('source_url', '') if args.samples_manifest else ''),
        'duration_seconds': duration, 'width': width, 'height': height,
        'interval_seconds': args.interval, 'sampling': 'all fixed-grid samples; no MAE prefilter',
        'sampling_command': command, 'sample_count': len(all_paths), 'hash_mode': args.hash_mode,
        'stability_eligible_samples': len(paths), 'stability_rejected_samples': len(all_paths)-len(paths),
        'stability': {key: stability[key] for key in ('status', 'threshold', 'width', 'height',
            'unique_adjacent_comparisons', 'rule', 'scikit_image_version', 'timings')},
        'stability_implementation_sha256': sha256(Path(__file__).with_name('stability_filter.py')),
        'algorithm': '64-bit DCT-II pHash: gray 32x32 LANCZOS; top-left 8x8; > median',
        'hash_bits': HASH_BITS, 'hamming_threshold': args.hash_distance,
        'duplicate_rule': 'first matching retained representative; discarded frames never become representatives',
        'retained_pages': len(pages), 'automatic_merges': len(samples)-len(pages),
        'ai_decisions_used': False, 'semantic_postprocessing': False,
        'dependencies': {'numpy': np.__version__, 'Pillow': PIL.__version__,
                         'reportlab': reportlab_version, 'pypdf': pypdf.__version__},
        'limitations': ['Fixed sampling can miss pages between grid points.',
                       'Both original neighbours are required; endpoints and short or changing real pages can be rejected.',
                       'Three qualifying sample points do not prove continuous stability between them.',
                       'Larger hash thresholds can merge different content; no threshold guarantees semantic equality.',
                       'Post-filter SSIM cannot find content already discarded by hashing.'], **provenance}
    if args.hash_mode == 'phash-dhash':
        metadata.update(dhash_algorithm='64 horizontal gradients: gray 9x8 LANCZOS; right > left; row-major',
            dhash_bits=HASH_BITS, dhash_hamming_threshold=args.dhash_distance,
            duplicate_rule='both inclusive thresholds must match the same first retained representative; no chaining')
    metadata['timings'] = {'stability': stability['timings'], 'hashes': hash_timings,
                           'comparison_seconds': comparison_seconds}
    hash_samples, hash_pages = samples, pages
    metadata['hash_retained_pages'] = len(hash_pages)
    source_hash_samples, source_hash_pages = restore_source_ids(hash_samples, hash_pages, eligible_ids,
                                                               args.interval, stability['samples'])
    save_json(output/'hash-samples.json', {**metadata, 'stage': 'hash_filter', 'samples': source_hash_samples})
    save_json(output/'hash-pages.json', {**metadata, 'stage': 'hash_filter', 'pages': source_hash_pages})
    save_json(output/'stage-status.json', {'stage': 'hash_complete', 'sample_count': len(all_paths),
                                         'stability_eligible_samples': len(paths),
                                         'hash_retained_pages': len(hash_pages)})
    if args.skip_ssim_audit:
        audit = {'status': 'skipped', 'reason': 'explicit --skip-ssim-audit',
                 'audit_only': True, 'automatic_deletions': 0}
        save_json(output/'ssim-audit.json', audit)
        ssim_assignments = [(i, None) for i in range(len(hash_pages))]
    else:
        print(json.dumps({'stage': 'retained_only_ssim_audit', 'pages': len(pages),
                          'unique_pairs': len(pages)*(len(pages)-1)//2}), flush=True)
        full_audit = audit_ssim(output, source_hash_pages, args.ssim_width, args.ssim_height, args.ssim_review_threshold,
                               args.ssim_threshold if args.ssim_mode == 'dedup' else None)
        audit = {key: full_audit[key] for key in ('status', 'audit_only', 'automatic_deletions',
                  'shape', 'unique_comparisons', 'width', 'height', 'timings',
                  'symmetric', 'unit_diagonal', 'scikit_image_version')}
        audit['report'] = str(output/'ssim-audit.html')
        audit['implementation_sha256'] = sha256(Path(__file__).with_name('ssim_audit.py'))
        if args.ssim_mode == 'dedup':
            ssim_assignments = [(row['representative_input_index'], row['ssim'])
                                for row in full_audit['selection']['assignments']]
        else:
            ssim_assignments = [(i, 1.0) for i in range(len(hash_pages))]
    samples, pages, decisions = compose_stages(paths, hashes, hash_samples, hash_pages,
        ssim_assignments, args.interval, dhashes if args.hash_mode == 'phash-dhash' else None)
    samples, pages = restore_source_ids(samples, pages, eligible_ids, args.interval, stability['samples'])
    decisions = restore_decisions(decisions, eligible_ids, args.interval)
    metadata.update(retained_pages=len(pages), automatic_merges=len(paths)-len(pages),
        ssim_mode='skipped' if args.skip_ssim_audit else args.ssim_mode,
        ssim_threshold=args.ssim_threshold, ssim_merges=len(hash_pages)-len(pages),
        final_distance_note='Distances to final representative may exceed hash thresholds; hash-stage distances are stored separately.')
    metadata['hash_duplicate_rule'] = metadata['duplicate_rule']
    if not args.skip_ssim_audit and args.ssim_mode == 'dedup':
        metadata['duplicate_rule'] = 'hash-stage first representative, then SSIM >= threshold to the first retained hash representative; neither stage uses transitive closure'
    save_json(output/'samples.json', {**metadata, 'samples': samples})
    save_json(output/'pages.json', {**metadata, 'pages': pages})
    save_json(output/'ssim-decisions.json', {'input_pages': len(hash_pages), 'retained_pages': len(pages),
        'mode': metadata['ssim_mode'], 'threshold': args.ssim_threshold, 'decisions': decisions,
        'merges': [decision for decision in decisions if not decision['retained']]})
    save_json(output/'stage-status.json', {'stage': 'ssim_complete', 'sample_count': len(all_paths),
        'stability_eligible_samples': len(paths),
        'hash_retained_pages': len(hash_pages), 'final_pages': len(pages), 'ssim_merges': len(hash_pages)-len(pages)})
    print(json.dumps({'stage': 'ssim_filter_complete', 'input_pages': len(hash_pages),
                      'retained_pages': len(pages), 'ssim_merges': len(hash_pages)-len(pages)}), flush=True)
    metadata['ssim_audit'] = audit
    started = time.perf_counter()
    exports, validations = {}, {}
    if args.output_format in ('pdf', 'both'):
        pdf_path = output/'slides.pdf'
        export_pdf(pdf_path, pages, width, height, args.title or video.stem, metadata['source_url'])
        validations['pdf'] = verify_pdf(pdf_path, pages)
        exports.update(pdf=str(pdf_path), pdf_sha256=sha256(pdf_path))
    if args.output_format in ('pptx', 'both'):
        pptx_path = output/'slides.pptx'
        export_pptx(pptx_path, pages, width, height, args.title or video.stem)
        validations['pptx'] = verify_pptx(pptx_path, pages)
        exports.update(pptx=str(pptx_path), pptx_sha256=sha256(pptx_path))
    metadata['timings']['export_and_verification_seconds'] = time.perf_counter()-started
    metadata['timings']['total_seconds'] = time.perf_counter()-started_all
    validation = {'page_count': len(pages), 'ordered_source_pixels_preserved': True, 'exports': validations}
    save_json(output/'run.json', {**metadata, 'argv': sys.argv, **exports, 'validation': validation,
        'completed_at_utc': datetime.now(timezone.utc).isoformat()})
    save_json(output/'stage-status.json', {'stage': 'completed', 'sample_count': len(samples),
        'stability_eligible_samples': len(paths), 'stability_rejected_samples': len(all_paths)-len(paths),
        'hash_retained_pages': len(hash_pages), 'final_pages': len(pages), 'ssim_merges': len(hash_pages)-len(pages)})
    print(json.dumps({'output': str(output), 'samples': len(samples), 'retained_pages': len(pages),
                      'stability_eligible_samples': len(paths),
                      'automatic_merges': len(paths)-len(pages), 'ssim_audit': audit['status'],
                      'timings': metadata['timings'], **exports}))
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except (ValueError, KeyError, OSError, subprocess.CalledProcessError) as error:
        sys.exit(f'Stopped: {error}. Existing output preserved.')
